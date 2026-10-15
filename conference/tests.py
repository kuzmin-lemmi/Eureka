from io import BytesIO, StringIO
import os
import shutil
import tempfile
from datetime import timedelta
from pathlib import Path

from django.contrib.auth.models import Group, Permission, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from .apps import ORGANIZERS_GROUP
from .models import ConferenceSettings, Document, FaqItem, NewsPost, ReviewLink, Section, Submission

MEDIA_ROOT = tempfile.mkdtemp()


def make_snils(number):
    """Корректный СНИЛС с контрольным числом для 9-значного номера."""
    digits = f"{number:09d}"
    total = sum(int(d) * w for d, w in zip(digits, range(9, 0, -1)))
    control = total if total < 100 else (0 if total in (100, 101) else (total % 101) % 100)
    return f"{digits[:3]}-{digits[3:6]}-{digits[6:]} {control:02d}"


@override_settings(MEDIA_ROOT=MEDIA_ROOT, MAX_UPLOAD_MB=1)
class SubmissionTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA_ROOT, ignore_errors=True)

    def setUp(self):
        self.section = Section.objects.create(name="Естественные науки")
        self.posts = 0

    def post(self, **overrides):
        # Каждая заявка по умолчанию от нового человека: одному СНИЛС нельзя дважды в одно направление.
        self.posts += 1
        data = {
            "email": "pupil@example.ru",
            "district": "Ленинский район",
            "school": "Гимназия № 87",
            "last_name": "Иванова",
            "first_name": "Мария",
            "grade": "9",
            "snils": make_snils(112233444 + self.posts),
            "title": "Маятник Фуко в домашних условиях",
            "section_slot": f"{self.section.pk}:1",
            "supervisor_last_name": "Петрова",
            "supervisor_first_name": "Анна",
            "supervisor_position": "учитель физики",
            "supervisor_phone": "+7 900 123-45-67",
            "file": SimpleUploadedFile("работа.docx", b"PK test", content_type="application/octet-stream"),
            "application_file": SimpleUploadedFile("заявка.pdf", b"%PDF-1.4"),
            "consent_file": SimpleUploadedFile("IMG_0001.JPG", b"JPEG test"),
            "consent": "on",
            **overrides,
        }
        return self.client.post(reverse("conference:submit"), data)

    def test_valid_submission_is_saved(self):
        response = self.post()
        self.assertRedirects(response, reverse("conference:thanks"))
        submission = Submission.objects.get()
        self.assertEqual(submission.admission, Submission.Admission.PENDING)
        self.assertEqual(submission.grade, 9)
        self.assertEqual(submission.original_filename, "работа.docx")
        self.assertNotIn("работа", submission.file.name)
        self.assertEqual(submission.coauthor, "")
        self.assertEqual(submission.application_original_filename, "заявка.pdf")
        self.assertEqual(submission.consent_original_filename, "IMG_0001.JPG")
        self.assertTrue(submission.consent_file.name.startswith("documents/"))

        thanks = self.client.get(reverse("conference:thanks"))
        self.assertContains(thanks, "Работа отправлена!")
        self.assertNotContains(thanks, submission.number)  # номер — для оргкомитета, ученику не нужен

    def test_presentation_rejected(self):
        self.post(file=SimpleUploadedFile("slides.pptx", b"PK"))
        self.assertFalse(Submission.objects.exists())

    def test_video_rejected(self):
        response = self.post(file=SimpleUploadedFile("video.mp4", b"\x00"))
        self.assertFalse(Submission.objects.exists())
        self.assertContains(response, "Недопустимый формат")

    def test_too_large_file_rejected(self):
        response = self.post(file=SimpleUploadedFile("big.pdf", b"0" * (1024 * 1024 + 1)))
        self.assertFalse(Submission.objects.exists())
        self.assertContains(response, "Файл больше 1 МБ")

    def test_documents_required(self):
        self.post(application_file="", consent_file="")
        self.assertFalse(Submission.objects.exists())

    def test_document_must_be_scan(self):
        response = self.post(consent_file=SimpleUploadedFile("согласие.docx", b"PK"))
        self.assertFalse(Submission.objects.exists())
        self.assertContains(response, "Загрузите согласие в PDF, JPG, PNG")

    def test_more_prefills_school_and_supervisor(self):
        self.post()
        response = self.client.get(reverse("conference:submit") + "?more=1")
        self.assertContains(response, 'value="Гимназия № 87"')
        self.assertContains(response, 'value="Петрова"')
        self.assertNotContains(response, 'value="Иванова"')
        # Без ?more анкета пустая.
        self.assertNotContains(self.client.get(reverse("conference:submit")), 'value="Петрова"')

    def test_only_first_group_is_offered_until_it_is_full(self):
        self.section.groups = 3
        self.section.capacity = 1
        self.section.save()
        page = self.client.get(reverse("conference:submit")).content.decode()
        self.assertIn("Естественные науки 1 — свободно 1 место", page)
        self.assertNotIn("Естественные науки 2", page)
        self.assertNotIn("Естественные науки 3", page)

        # Чужую секцию из будущего подставить нельзя.
        self.post(section_slot=f"{self.section.pk}:2")
        self.assertFalse(Submission.objects.exists())

        self.post()
        page = self.client.get(reverse("conference:submit")).content.decode()
        self.assertIn("Естественные науки 2 — свободно 1 место", page)
        self.assertNotIn("Естественные науки 1", page)

        self.post(last_name="Петров", section_slot=f"{self.section.pk}:2")
        submission = Submission.objects.get(last_name="Петров")
        self.assertEqual((submission.section_number, submission.section_label), (2, "Естественные науки 2"))

    def test_when_all_groups_are_full_last_one_is_shown_disabled(self):
        self.section.groups = 2
        self.section.capacity = 1
        self.section.save()
        self.post()
        self.post(last_name="Петров", section_slot=f"{self.section.pk}:2")
        page = self.client.get(reverse("conference:submit")).content.decode()
        self.assertIn("Естественные науки 2 — мест нет", page)
        self.assertNotIn("Естественные науки 1", page)

    def test_group_number_out_of_range_rejected(self):
        self.post(section_slot=f"{self.section.pk}:2")
        self.assertFalse(Submission.objects.exists())

    def test_full_section_is_disabled_and_rejected(self):
        self.section.capacity = 1
        self.section.save()
        self.post()
        self.assertEqual(Submission.objects.count(), 1)

        response = self.client.get(reverse("conference:submit"))
        self.assertContains(response, "Естественные науки — мест нет")
        self.assertContains(response, "disabled")

        response = self.post(last_name="Петров")
        self.assertContains(response, "Эта секция уже заполнена")
        self.assertEqual(Submission.objects.count(), 1)

    def test_work_with_coauthor_takes_one_place(self):
        self.section.capacity = 2
        self.section.save()
        self.post(coauthor_last_name="Сидоров", coauthor_first_name="Пётр", coauthor_grade="8", coauthor_snils="123-456-789 64")
        self.assertEqual(self.section.free_places(), {1: 1})
        self.post(last_name="Петров", coauthor_last_name="Орлов", coauthor_first_name="Иван", coauthor_grade="8",
                  coauthor_snils="123-456-789 64")
        self.assertEqual(Submission.objects.count(), 2)
        self.assertEqual(self.section.free_places(), {1: 0})

    def test_deleted_submission_frees_place(self):
        self.section.capacity = 1
        self.section.save()
        self.post()
        Submission.objects.get().delete()
        self.post(last_name="Петров")
        self.assertEqual(Submission.objects.count(), 1)
        self.post(last_name="Сидоров")  # место снова занято
        self.assertEqual(Submission.objects.count(), 1)

    def test_grade_out_of_range_rejected(self):
        for grade in ("7", "12"):
            self.post(grade=grade)
        self.assertFalse(Submission.objects.exists())

    def test_partial_coauthor_rejected(self):
        response = self.post(coauthor_last_name="Сидоров")
        self.assertFalse(Submission.objects.exists())
        self.assertContains(response, "Заполните, если у работы есть соавтор.", count=3)

    def test_full_coauthor_saved(self):
        self.post(coauthor_last_name="Сидоров", coauthor_first_name="Пётр", coauthor_grade="8", coauthor_snils="123-456-789 64")
        self.assertEqual(Submission.objects.get().coauthor, "Сидоров Пётр, 8 класс")

    def test_supervisor_middle_name_optional(self):
        self.post(supervisor_middle_name="")
        self.assertEqual(Submission.objects.get().supervisor, "Петрова Анна")

    def test_snils_is_required_and_normalized(self):
        self.post(snils="11223344595")
        self.assertEqual(Submission.objects.get().snils, "112-233-445 95")

    def test_missing_snils_rejected(self):
        self.post(snils="")
        self.assertFalse(Submission.objects.exists())

    def test_snils_typo_rejected(self):
        response = self.post(snils="112-233-445 96")
        self.assertFalse(Submission.objects.exists())
        self.assertContains(response, "опечатка")

    def test_short_snils_rejected(self):
        self.post(snils="112-233")
        self.assertFalse(Submission.objects.exists())

    def test_coauthor_requires_snils(self):
        response = self.post(coauthor_last_name="Сидоров", coauthor_first_name="Пётр", coauthor_grade="8")
        self.assertFalse(Submission.objects.exists())
        self.assertContains(response, "Заполните, если у работы есть соавтор.")

    def test_coauthor_snils_saved(self):
        self.post(
            coauthor_last_name="Сидоров", coauthor_first_name="Пётр", coauthor_grade="8", coauthor_snils="12345678964"
        )
        self.assertEqual(Submission.objects.get().coauthor_snils, "123-456-789 64")

    def test_second_supervisor_is_optional(self):
        self.post()
        submission = Submission.objects.get()
        self.assertEqual(submission.supervisor2, "")

    def test_second_supervisor_saved(self):
        self.post(
            supervisor2_last_name="Сидорова",
            supervisor2_first_name="Мария",
            supervisor2_middle_name="Ивановна",
            supervisor2_position="учитель биологии",
            supervisor2_phone="+7 900 765-43-21",
        )
        self.assertEqual(Submission.objects.get().supervisor2, "Сидорова Мария Ивановна")

    def test_partial_second_supervisor_rejected(self):
        response = self.post(supervisor2_last_name="Сидорова")
        self.assertFalse(Submission.objects.exists())
        self.assertContains(response, "Заполните, если у работы два руководителя.", count=3)

    def test_bad_second_supervisor_phone_rejected(self):
        self.post(
            supervisor2_last_name="Сидорова",
            supervisor2_first_name="Мария",
            supervisor2_position="учитель",
            supervisor2_phone="123",
        )
        self.assertFalse(Submission.objects.exists())

    def test_second_supervisor_block_open_on_error(self):
        response = self.post(supervisor2_last_name="Сидорова")
        self.assertContains(response, 'id="step-supervisor2" open')

    def test_bad_supervisor_phone_rejected(self):
        self.post(supervisor_phone="123")
        self.assertFalse(Submission.objects.exists())

    def test_consent_required(self):
        self.post(consent="")
        self.assertFalse(Submission.objects.exists())

    def test_inactive_section_not_allowed(self):
        hidden = Section.objects.create(name="Скрытая", is_active=False)
        self.post(section_slot=f"{hidden.pk}:1")
        self.assertFalse(Submission.objects.exists())

    def test_closed_after_deadline(self):
        conf = ConferenceSettings.load()
        conf.submissions_deadline = timezone.now() - timedelta(minutes=1)
        conf.save()
        response = self.post()
        self.assertContains(response, "Приём работ закрыт")
        self.assertFalse(Submission.objects.exists())

    def test_honeypot_discards_spam(self):
        response = self.post(website="http://spam.example")
        self.assertRedirects(response, reverse("conference:thanks"))
        self.assertFalse(Submission.objects.exists())

    def test_file_download_requires_staff(self):
        self.post()
        url = reverse("conference:download_submission", args=[Submission.objects.get().pk])

        self.assertEqual(self.client.get(url).status_code, 302)

        staff = User.objects.create_user("org", password="pass", is_staff=True)
        self.client.force_login(staff)
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])

        pk = Submission.objects.get().pk
        consent_url = reverse("conference:download_submission_file", args=[pk, "consent"])
        response = self.client.get(consent_url)
        self.assertEqual(response.status_code, 200)
        self.assertIn(f"{pk:04d}", response["Content-Disposition"])
        self.assertEqual(self.client.get(reverse("conference:download_submission_file", args=[pk, "x"])).status_code, 404)

        self.client.logout()
        self.assertEqual(self.client.get(consent_url).status_code, 302)

    def test_blank_is_public(self):
        self.assertEqual(self.client.get(reverse("conference:blank", args=["consent"])).status_code, 404)
        conf = ConferenceSettings.load()
        conf.consent_pdf = SimpleUploadedFile("согласие.pdf", b"%PDF blank")
        conf.save()
        response = self.client.get(reverse("conference:blank", args=["consent"]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertContains(self.client.get(reverse("conference:submit")), reverse("conference:blank", args=["consent"]))


class OrganizersTests(TestCase):
    def test_organizer_can_manage_submissions_but_not_users(self):
        organizer = User.objects.create_user("org", password="pass", is_staff=True)
        organizer.groups.add(Group.objects.get(name=ORGANIZERS_GROUP))
        self.client.force_login(organizer)

        self.assertEqual(self.client.get(reverse("admin:conference_submission_changelist")).status_code, 200)
        self.assertEqual(self.client.get(reverse("admin:conference_section_add")).status_code, 200)
        self.assertEqual(self.client.get(reverse("admin:auth_user_changelist")).status_code, 403)


class LoadSectionsTests(TestCase):
    def test_load_sections_creates_list_and_replaces_demo(self):
        from django.core.management import call_command

        Section.objects.create(name="Демо")
        call_command("load_sections", "--replace", stdout=open(os.devnull, "w"))
        self.assertFalse(Section.objects.filter(name="Демо").exists())
        self.assertEqual(Section.objects.get(name="Региональное краеведение").groups, 3)
        self.assertEqual(Section.objects.count(), 13)


class PagesTests(TestCase):
    def test_fresh_news_banner(self):
        old = NewsPost.objects.create(title="Старая новость", text="…", published_at=timezone.now() - timedelta(days=8))
        self.assertNotContains(self.client.get(reverse("conference:home")), "news-flash")

        fresh = NewsPost.objects.create(title="Перенесли дедлайн", text="…")
        response = self.client.get(reverse("conference:home"))
        self.assertContains(response, f'class="news-flash" href="{reverse("conference:news_detail", args=[fresh.pk])}"')
        self.assertContains(response, f'id="news-{fresh.pk}"')
        self.assertContains(response, f'id="news-{old.pk}"')

    def test_home_block_order(self):
        Section.objects.create(name="Биология")
        html = self.client.get(reverse("conference:home")).content.decode()
        positions = [html.index(f'id="{block}"') for block in ("how", "documents", "sections", "news")]
        self.assertEqual(positions, sorted(positions))

    def test_home_shows_active_sections_only(self):
        Section.objects.create(name="Биология")
        Section.objects.create(name="Секретная", is_active=False)
        response = self.client.get(reverse("conference:home"))
        self.assertContains(response, "Биология")
        self.assertNotContains(response, "Секретная")

    def test_settings_admin_opens_single_record(self):
        admin = User.objects.create_superuser("admin", password="pass")
        self.client.force_login(admin)
        response = self.client.get(reverse("admin:conference_conferencesettings_changelist"))
        self.assertRedirects(response, reverse("admin:conference_conferencesettings_change", args=[1]))


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class NewsAndDocumentsTests(TestCase):
    def home(self):
        return self.client.get(reverse("conference:home"))

    def test_home_shows_only_published_and_past_news(self):
        NewsPost.objects.create(title="Открыта", text="Текст")
        NewsPost.objects.create(title="Черновик", text="Текст", is_published=False)
        NewsPost.objects.create(title="Завтра", text="Текст", published_at=timezone.now() + timedelta(days=1))
        response = self.home()
        self.assertContains(response, "Открыта")
        self.assertNotContains(response, "Черновик")
        self.assertNotContains(response, "Завтра")

    def test_news_text_is_on_home(self):
        NewsPost.objects.create(title="Итоги", text="Первый абзац\n\nВторой абзац")
        self.assertContains(self.home(), "Второй абзац")

    def test_home_shows_three_latest_news_and_list_has_all(self):
        for i in range(12):
            NewsPost.objects.create(title=f"Новость {i:02d}", text="…", published_at=timezone.now() - timedelta(days=i))
        home = self.home()
        self.assertEqual(home.content.decode().count('class="card news-card'), 3)
        self.assertContains(home, "Новость 00")
        self.assertNotContains(home, "Новость 03")
        self.assertContains(home, "Все новости")
        page = self.client.get(reverse("conference:news"))
        self.assertEqual(page.content.decode().count('class="card news-card'), 10)
        self.assertContains(page, "Страница 1 из 2")
        self.assertContains(self.client.get(reverse("conference:news") + "?page=2"), "Новость 11")

    def test_empty_blocks_have_placeholder(self):
        response = self.home()
        self.assertContains(response, "Новостей пока нет")
        self.assertContains(response, "Документы пока не опубликованы")

    def test_old_urls_redirect_to_home_anchors(self):
        self.assertRedirects(
            self.client.get(reverse("conference:documents")), reverse("conference:home") + "#documents"
        )

    def test_document_is_listed_and_downloadable(self):
        doc = Document.objects.create(title="Положение", file=SimpleUploadedFile("pol.pdf", b"%PDF-1.4"))
        self.assertContains(self.home(), "Положение")
        response = self.client.get(reverse("conference:download_document", args=[doc.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])

    def test_unpublished_document_hidden(self):
        doc = Document.objects.create(title="Секрет", file=SimpleUploadedFile("s.pdf", b"x"), is_published=False)
        self.assertNotContains(self.home(), "Секрет")
        self.assertEqual(self.client.get(reverse("conference:download_document", args=[doc.pk])).status_code, 404)

    def test_organizer_can_publish_news_and_documents(self):
        organizer = User.objects.create_user("org2", password="pass", is_staff=True)
        organizer.groups.add(Group.objects.get(name=ORGANIZERS_GROUP))
        self.client.force_login(organizer)
        self.assertEqual(self.client.get(reverse("admin:conference_newspost_add")).status_code, 200)
        self.assertEqual(self.client.get(reverse("admin:conference_document_add")).status_code, 200)


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class BlanksTests(TestCase):
    def test_load_blanks_command_fills_all_templates(self):
        from django.core.management import call_command

        call_command("load_blanks", stdout=StringIO())
        conf = ConferenceSettings.load()
        for field in ("application_pdf", "application_template", "consent_pdf", "consent_template", "requirements_template"):
            self.assertTrue(getattr(conf, field), field)

    def test_blanks_are_listed_on_home_and_downloadable(self):
        from django.core.management import call_command

        call_command("load_blanks", stdout=StringIO())
        home = self.client.get(reverse("conference:home"))
        for title in ("Заявка на участие", "Согласие родителя (законного представителя)", "Требования к работе"):
            self.assertContains(home, title)
        for kind in ("application", "application-word", "consent", "consent-word", "requirements"):
            response = self.client.get(reverse("conference:blank", args=[kind]))
            self.assertEqual(response.status_code, 200)
            self.assertIn("attachment", response["Content-Disposition"])

    def test_anketa_links_requirements_and_blanks(self):
        from django.core.management import call_command

        Section.objects.create(name="Физика")
        call_command("load_blanks", stdout=StringIO())
        page = self.client.get(reverse("conference:submit"))
        for kind in ("application", "application-word", "consent", "consent-word", "requirements"):
            self.assertContains(page, reverse("conference:blank", args=[kind]))

    def test_home_without_blanks_hides_them(self):
        self.assertNotContains(self.client.get(reverse("conference:home")), "Требования к работе")


class SectionCapacityTests(TestCase):
    def test_new_section_has_12_places(self):
        section = Section.objects.create(name="Физика")
        self.assertEqual(section.capacity, 12)
        self.assertEqual(sum(section.free_places().values()), 12)


@override_settings(MEDIA_ROOT=MEDIA_ROOT, MAX_UPLOAD_MB=1)
class ParticipantAccessTests(TestCase):
    SNILS = "112-233-445 95"

    def setUp(self):
        self.section = Section.objects.create(name="Экология", groups=2, capacity=3)
        self.other = Section.objects.create(name="Краеведение")

    def data(self, **overrides):
        return {
            "email": "Pupil@Example.ru",
            "district": "Ленинский район",
            "school": "Гимназия № 87",
            "last_name": "Иванова",
            "first_name": "Мария",
            "grade": "9",
            "snils": self.SNILS,
            "title": "Маятник Фуко",
            "section_slot": f"{self.section.pk}:1",
            "supervisor_last_name": "Петрова",
            "supervisor_first_name": "Анна",
            "supervisor_position": "учитель физики",
            "supervisor_phone": "+7 900 123-45-67",
            "file": SimpleUploadedFile("работа.docx", b"PK test"),
            "application_file": SimpleUploadedFile("заявка.pdf", b"%PDF-1.4"),
            "consent_file": SimpleUploadedFile("согласие.jpg", b"JPEG"),
            "consent": "on",
            **overrides,
        }

    def submit(self, **overrides):
        return self.client.post(reverse("conference:submit"), self.data(**overrides))

    def login(self, snils=SNILS, email="pupil@example.ru"):
        return self.client.post(reverse("conference:my"), {"snils": snils, "email": email})

    def test_second_submission_to_same_direction_rejected(self):
        self.submit()
        response = self.submit(section_slot=f"{self.section.pk}:1", title="Другая работа")
        self.assertContains(response, "уже есть заявка в направлении «Экология»")
        self.assertEqual(Submission.objects.count(), 1)

    def test_same_person_can_apply_to_other_direction(self):
        self.submit()
        self.submit(section_slot=f"{self.other.pk}:1", title="Другая работа")
        self.assertEqual(Submission.objects.count(), 2)

    def test_deleted_submission_allows_new_one(self):
        self.submit()
        self.client.post(reverse("conference:my_delete", args=[Submission.objects.get().pk]))
        self.submit()
        self.assertEqual(Submission.objects.count(), 1)

    def test_coauthor_is_not_restricted(self):
        self.submit()
        self.submit(snils=make_snils(200000001), coauthor_last_name="Иванова", coauthor_first_name="Мария",
                    coauthor_grade="9", coauthor_snils=self.SNILS)
        self.assertEqual(Submission.objects.count(), 2)

    def test_database_blocks_duplicate(self):
        from django.db import IntegrityError, transaction

        self.submit()
        first = Submission.objects.get()
        first.pk = None
        with self.assertRaises(IntegrityError), transaction.atomic():
            first.save()

    def test_after_submit_participant_sees_my_page(self):
        self.submit()
        response = self.client.get(reverse("conference:my"))
        self.assertContains(response, "Маятник Фуко")
        self.assertContains(response, "Изменить заявку")

    def test_login_requires_matching_snils_and_email(self):
        self.submit()
        self.client.post(reverse("conference:my_logout"))
        response = self.login(email="other@example.ru")
        self.assertContains(response, "Заявки не найдены")
        self.assertNotContains(response, "Маятник Фуко")

        response = self.login(snils="11223344595", email="PUPIL@example.ru")
        self.assertRedirects(response, reverse("conference:my"))
        self.assertContains(self.client.get(reverse("conference:my")), "Маятник Фуко")

    def test_too_many_failed_attempts_blocked(self):
        from django.core.cache import cache

        cache.clear()
        self.submit()
        self.client.post(reverse("conference:my_logout"))
        for _ in range(10):
            self.login(email="wrong@example.ru")
        response = self.login()
        self.assertContains(response, "Слишком много попыток")
        cache.clear()

    def test_edit_keeps_files_and_frees_own_place(self):
        self.section.capacity = 1
        self.section.save()
        self.submit()
        submission = Submission.objects.get()
        url = reverse("conference:my_edit", args=[submission.pk])
        data = self.data(title="Новое название")
        for name in ("file", "application_file", "consent_file", "consent"):
            data.pop(name)
        response = self.client.post(url, data)
        self.assertRedirects(response, reverse("conference:my"))
        submission.refresh_from_db()
        self.assertEqual(submission.title, "Новое название")
        self.assertEqual(submission.original_filename, "работа.docx")
        self.assertIsNotNone(submission.edited_by_participant_at)

    def test_my_page_has_no_status_but_shows_message(self):
        self.submit()
        Submission.objects.update(participant_message="Добавьте список литературы")
        response = self.client.get(reverse("conference:my"))
        self.assertContains(response, "Добавьте список литературы")
        self.assertContains(response, "Удалить")
        for word in ("Принята", "Отклонена", "Допущен"):
            self.assertNotContains(response, word)

    def test_participant_deletes_submission_with_files(self):
        self.submit()
        self.submit(section_slot=f"{self.other.pk}:1", title="Другая работа")
        submission = Submission.objects.get(section=self.section)
        paths = [Path(f.path) for f in (submission.file, submission.application_file, submission.consent_file)]
        self.assertTrue(all(p.exists() for p in paths))
        url = reverse("conference:my_delete", args=[submission.pk])
        self.assertContains(self.client.get(url), "Да, удалить заявку")
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(url)
        self.assertRedirects(response, reverse("conference:my"))
        self.assertFalse(Submission.objects.filter(pk=submission.pk).exists())
        self.assertFalse(any(p.exists() for p in paths))
        self.assertContains(self.client.get(reverse("conference:my")), "Другая работа")  # вход сохранился

    def test_deleting_last_submission_logs_out(self):
        self.submit()
        self.client.post(reverse("conference:my_delete", args=[Submission.objects.get().pk]))
        self.assertContains(self.client.get(reverse("conference:my")), "Найти заявки")

    def test_closed_intake_forbids_edit_and_delete(self):
        self.submit()
        conf = ConferenceSettings.load()
        conf.submissions_open = False
        conf.save()
        submission = Submission.objects.get()
        my = reverse("conference:my")
        self.assertRedirects(self.client.get(reverse("conference:my_edit", args=[submission.pk])), my)
        self.assertRedirects(self.client.post(reverse("conference:my_delete", args=[submission.pk])), my)
        self.assertTrue(Submission.objects.exists())
        response = self.client.get(my)
        self.assertContains(response, "Приём работ закрыт")
        self.assertNotContains(response, "Удалить")

    def test_cannot_delete_someone_elses_submission(self):
        self.submit()
        foreign = Submission.objects.get()
        self.client.post(reverse("conference:my_logout"))
        self.submit(snils=make_snils(200000001), email="other@example.ru", section_slot=f"{self.other.pk}:1")
        self.assertEqual(self.client.post(reverse("conference:my_delete", args=[foreign.pk])).status_code, 404)
        self.assertTrue(Submission.objects.filter(pk=foreign.pk).exists())

    def test_cannot_edit_someone_elses_submission(self):
        self.submit()
        foreign = Submission.objects.get()
        self.client.post(reverse("conference:my_logout"))
        self.submit(snils=make_snils(200000001), email="other@example.ru", section_slot=f"{self.other.pk}:1")
        response = self.client.get(reverse("conference:my_edit", args=[foreign.pk]))
        self.assertEqual(response.status_code, 404)

    def test_edit_without_login_redirects(self):
        self.submit()
        submission = Submission.objects.get()
        self.client.post(reverse("conference:my_logout"))
        response = self.client.get(reverse("conference:my_edit", args=[submission.pk]))
        self.assertRedirects(response, reverse("conference:my"))


@override_settings(MEDIA_ROOT=MEDIA_ROOT, MAX_UPLOAD_MB=1)
class SectionsExportTests(TestCase):
    def setUp(self):
        self.section = Section.objects.create(name="Искусство слова и лингвистика. Язык в современном мире", groups=2)
        for i, number in enumerate((1, 2)):
            Submission.objects.create(
                section=self.section, section_number=number, snils=make_snils(300000000 + i),
                last_name=f"Участник{i}", first_name="Имя", grade=9, email="p@example.ru",
                school="Гимназия № 87", district="Ленинский", title=f"Работа {i}",
                supervisor_last_name="Петрова", supervisor_first_name="Анна", supervisor_position="учитель",
                supervisor_phone="+79001234567", file="x.docx",
            )

    def download(self):
        from openpyxl import load_workbook

        response = self.client.get(reverse("admin:conference_submission_export_sections"))
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(b"".join(response.streaming_content)))

    def test_export_has_sheet_per_group(self):
        self.client.force_login(User.objects.create_superuser("admin", password="pass"))
        wb = self.download()
        self.assertEqual(len(wb.sheetnames), 3)
        self.assertTrue(all(len(name) <= 31 for name in wb.sheetnames))
        first = wb[wb.sheetnames[1]]
        names = [row[2] for row in first.iter_rows(min_row=4, values_only=True)]
        self.assertEqual(names, ["Участник0 Имя"])
        self.assertNotIn("300-000-000", str(list(first.values)))  # СНИЛС не выгружается
        self.assertEqual(wb["Сводка"]["C2"].value, 1)

    def test_export_requires_staff(self):
        response = self.client.get(reverse("admin:conference_submission_export_sections"))
        self.assertEqual(response.status_code, 302)

    def test_organizer_sees_export_button(self):
        organizer = User.objects.create_user("org", password="pass", is_staff=True)
        organizer.groups.add(Group.objects.get(name=ORGANIZERS_GROUP))
        self.client.force_login(organizer)
        response = self.client.get(reverse("admin:conference_submission_changelist"))
        self.assertContains(response, reverse("admin:conference_submission_exports"))
        self.download()

    def test_exports_page_has_four_stages_and_student_files(self):
        from docx import Document as Docx
        from openpyxl import load_workbook

        Submission.objects.update(coauthor_last_name="Соавтор", coauthor_first_name="Ира", coauthor_grade=8)
        self.client.force_login(User.objects.create_superuser("admin", password="pass"))
        page = self.client.get(reverse("admin:conference_submission_exports"))
        for n in range(1, 5):
            self.assertContains(page, f"Этап {n}")

        response = self.client.get(reverse("conference:office_students_excel"))
        ws = load_workbook(BytesIO(b"".join(response.streaming_content))).active
        title, _, *rows = list(ws.values)
        self.assertEqual(title[0], "Список учеников по секциям. XII межрегиональная научно-практическая конференция школьников «Эврика», 2027")
        self.assertEqual(rows[0][:4], ("№", "Секция", "ФИО", "Кл."))
        self.assertEqual(len(rows), 1 + 4)  # две заявки по два ученика
        self.assertEqual(rows[2][2], "Соавтор Ира")
        self.assertEqual(rows[2][3], 8)

        conf = ConferenceSettings.load()
        conf.edition, conf.year = "XIII", 2028
        conf.save()
        response = self.client.get(reverse("conference:office_students_word"))
        self.assertIn("XIII межрегиональная", Docx(BytesIO(b"".join(response.streaming_content))).paragraphs[0].text)
        response = self.client.get(reverse("admin:conference_submission_export_sections"))
        sheet = load_workbook(BytesIO(b"".join(response.streaming_content))).worksheets[1]
        self.assertTrue(sheet["A2"].value.startswith("XIII межрегиональная научно-практическая конференция школьников «Эврика», 2028"))

        response = self.client.get(reverse("conference:office_students_word"))
        doc = Docx(BytesIO(b"".join(response.streaming_content)))
        self.assertEqual(len(doc.tables[0].rows), 1 + 4)


@override_settings(MEDIA_ROOT=MEDIA_ROOT, MAX_UPLOAD_MB=1)
class OfficeTests(TestCase):
    def setUp(self):
        from django.core.cache import cache
        from django.core.files.base import ContentFile

        cache.clear()
        self.section = Section.objects.create(name="Здоровое поколение", groups=2, capacity=3)
        self.other = Section.objects.create(name="Краеведение")
        for i, (section, number) in enumerate(((self.section, 1), (self.section, 1), (self.section, 2))):
            s = Submission(
                section=section, section_number=number, snils=make_snils(400000000 + i),
                last_name=f"Фамилия{i}", first_name="Имя", grade=9, email=f"p{i}@example.ru",
                school="Гимназия № 87", district="Ленинский", title=f"Работа номер {i}",
                supervisor_last_name="Петрова", supervisor_first_name="Анна", supervisor_position="учитель",
                supervisor_phone="+79001234567",
            )
            s.file.save("работа.docx", ContentFile(f"содержимое {i}".encode()), save=False)
            s.original_filename = "работа.docx"
            s.save()
        self.organizer = User.objects.create_user("olga", password="Pass-1234!", is_staff=True)
        self.organizer.groups.add(Group.objects.get(name=ORGANIZERS_GROUP))

    def test_footer_has_office_login_link(self):
        self.assertContains(self.client.get(reverse("conference:home")), reverse("conference:office_login"))

    def test_office_requires_login(self):
        response = self.client.get(reverse("conference:office"))
        self.assertRedirects(response, reverse("conference:office_login") + "?next=/office/")

    def test_login_and_dashboard_counts(self):
        response = self.client.post(reverse("conference:office_login"), {"username": "olga", "password": "Pass-1234!"})
        self.assertRedirects(response, reverse("conference:office"))
        response = self.client.get(reverse("conference:office"))
        self.assertContains(response, "Здоровое поколение 1")
        self.assertContains(response, "<span>2/3</span>", html=False)
        self.assertContains(response, "<span>1/3</span>", html=False)
        self.assertContains(response, "<b>3</b> работ подано", html=False)

    def test_wrong_password_and_lockout(self):
        for _ in range(10):
            response = self.client.post(reverse("conference:office_login"), {"username": "olga", "password": "нет"})
        self.assertContains(response, "Неверный логин или пароль")
        response = self.client.post(reverse("conference:office_login"), {"username": "olga", "password": "Pass-1234!"})
        self.assertContains(response, "Слишком много неудачных попыток")

    def test_non_staff_cannot_enter(self):
        User.objects.create_user("pupil", password="Pass-1234!")
        response = self.client.post(reverse("conference:office_login"), {"username": "pupil", "password": "Pass-1234!"})
        self.assertContains(response, "нет доступа к кабинету")

    def test_organizer_closes_and_opens_intake(self):
        self.client.force_login(self.organizer)
        self.assertContains(self.client.get(reverse("conference:office")), "Закрыть приём")
        self.client.post(reverse("conference:office_intake"))
        self.assertFalse(ConferenceSettings.load().is_submission_open)
        self.assertContains(self.client.get(reverse("conference:office")), "Открыть приём")
        self.assertTrue(self.client.get(reverse("conference:submit")).context["closed"])

        conf = ConferenceSettings.load()
        conf.submissions_deadline = timezone.now() - timedelta(days=1)
        conf.save()
        self.client.post(reverse("conference:office_intake"), {"open": "1"})
        conf = ConferenceSettings.load()
        self.assertTrue(conf.is_submission_open)
        self.assertIsNone(conf.submissions_deadline)  # прошедший срок снят, иначе приём остался бы закрытым

    def test_admission_only_after_intake_closed(self):
        self.client.force_login(self.organizer)
        section_url = reverse("conference:office_section", args=[self.section.pk, 1])
        submission = Submission.objects.get(last_name="Фамилия0")
        url = reverse("conference:office_admission", args=[submission.pk])
        self.assertNotContains(self.client.get(section_url), "Не допущен</button>")
        self.client.post(url, {"admission": "admitted"})
        submission.refresh_from_db()
        self.assertEqual(submission.admission, "")

        conf = ConferenceSettings.load()
        conf.submissions_open = False
        conf.save()
        self.assertContains(self.client.get(section_url), "Не допущен</button>")
        response = self.client.post(url, {"admission": "admitted"})
        self.assertRedirects(response, f"{section_url}#s{submission.pk}", fetch_redirect_response=False)
        submission.refresh_from_db()
        self.assertEqual(submission.admission, "admitted")
        self.assertContains(self.client.get(section_url), 'допущено <span data-count="admitted">1</span>', html=False)
        self.client.post(url, {"admission": "not_admitted"})
        submission.refresh_from_db()
        self.assertEqual(submission.admission, "not_admitted")
        response = self.client.post(url, {"admission": "not_admitted"}, headers={"Accept": "application/json"})
        self.assertEqual(response.json(), {"admission": "not_admitted", "award": "", "admitted": 0, "not_admitted": 1, "pending": 1})
        self.client.post(url, {"admission": ""})  # «Снять отметку»
        submission.refresh_from_db()
        self.assertEqual(submission.admission, "")
        self.assertEqual(self.client.post(url, {"admission": "maybe"}).status_code, 404)

    def test_award_only_for_admitted_and_cleared_with_admission(self):
        self.client.force_login(self.organizer)
        conf = ConferenceSettings.load()
        conf.submissions_open = False
        conf.save()
        submission = Submission.objects.get(last_name="Фамилия0")
        award_url = reverse("conference:office_award", args=[submission.pk])
        self.client.post(award_url, {"award": "place1"})
        submission.refresh_from_db()
        self.assertEqual(submission.award, "")  # не допущен — итог не ставится

        self.client.post(reverse("conference:office_admission", args=[submission.pk]), {"admission": "admitted"})
        response = self.client.post(award_url, {"award": "social"}, headers={"Accept": "application/json"})
        self.assertEqual(response.json(), {"award": "social"})
        self.assertEqual(self.client.post(award_url, {"award": "place9"}).status_code, 404)
        self.assertContains(self.client.get(reverse("conference:office_section", args=[self.section.pk, 1])),
                            'value="social" selected', html=False)

        response = self.client.post(reverse("conference:office_admission", args=[submission.pk]), {"admission": ""},
                                    headers={"Accept": "application/json"})
        self.assertEqual(response.json()["award"], "")
        submission.refresh_from_db()
        self.assertEqual(submission.award, "")

    def test_winners_exports(self):
        from docx import Document
        from openpyxl import load_workbook

        Submission.objects.filter(last_name="Фамилия0").update(admission="admitted", award="social")
        Submission.objects.filter(last_name="Фамилия1").update(
            admission="admitted", award="place2", coauthor_last_name="Соавторова", coauthor_first_name="Ира", coauthor_grade=9,
        )
        Submission.objects.filter(last_name="Фамилия2").update(admission="admitted", award="place1")
        self.client.force_login(self.organizer)
        self.assertContains(self.client.get(reverse("conference:office")), "Отмечено работ: 3")

        response = self.client.get(reverse("conference:office_winners_excel"))
        rows = list(load_workbook(BytesIO(response.getvalue())).active.values)
        self.assertEqual(rows[0][0], "Победители и призеры XII межрегиональной научно-практической конференции «Эврика»")
        self.assertEqual(rows[2], ("№", "ФИО", "Название работы", "Секция", "ОУ", "Статус"))
        data = rows[3:]
        # Секции по алфавиту, внутри — места, затем номинации; соавтор в той же строке.
        self.assertEqual([(r[3], r[5]) for r in data], [
            ("Здоровое поколение 1", "2 место"),
            ("Здоровое поколение 1", "номинация: за изучение социально значимых объектов"),
            ("Здоровое поколение 2", "1 место"),
        ])
        self.assertEqual(data[0][1], "Фамилия1 Имя,\nСоавторова Ира")

        word = self.client.get(reverse("conference:office_winners_word"))
        self.assertEqual(len(Document(BytesIO(word.getvalue())).tables[0].rows), 1 + 3)
        pdf = self.client.get(reverse("conference:office_winners_pdf"))
        self.assertTrue(pdf.getvalue().startswith(b"%PDF"))

    def test_places_export_has_children_with_snils_and_only_places(self):
        from docx import Document
        from openpyxl import load_workbook

        Submission.objects.filter(last_name="Фамилия0").update(admission="admitted", award="social")  # номинация — не берём
        Submission.objects.filter(last_name="Фамилия1").update(
            admission="admitted", award="place2", coauthor_last_name="Соавторова", coauthor_first_name="Ира",
            coauthor_grade=9, coauthor_snils="112-233-445 95",
        )
        Submission.objects.filter(last_name="Фамилия2").update(admission="admitted", award="place1")
        self.client.force_login(self.organizer)
        self.assertContains(self.client.get(reverse("conference:office")), "Работ с местом: 2")

        response = self.client.get(reverse("conference:office_places_excel"))
        rows = list(load_workbook(BytesIO(response.getvalue())).active.values)
        self.assertEqual(rows[2], ("№", "ФИО", "Название работы", "Секция", "ОУ", "Статус", "СНИЛС"))
        data = rows[3:]
        self.assertEqual([(r[1], r[5]) for r in data], [
            ("Фамилия1 Имя", "2 место"), ("Соавторова Ира", "2 место"), ("Фамилия2 Имя", "1 место"),
        ])
        self.assertEqual(data[1][6], "112-233-445 95")  # СНИЛС соавтора, а не основного автора
        self.assertTrue(all(r[6] for r in data))

        self.assertEqual(len(Document(BytesIO(self.client.get(reverse("conference:office_places_word")).getvalue())).tables[0].rows), 1 + 3)
        self.assertTrue(self.client.get(reverse("conference:office_places_pdf")).getvalue().startswith(b"%PDF"))

    def test_settings_page_saves_fields_and_replaces_blank(self):
        self.client.force_login(self.organizer)
        url = reverse("conference:office_settings")
        self.assertContains(self.client.get(reverse("conference:office")), url)
        page = self.client.get(url)
        for text in ("Номер конференции", "Бланки документов", "Политика обработки персональных данных"):
            self.assertContains(page, text)

        data = {name: value for name, value in page.context["form"].initial.items() if value is not None and not hasattr(value, "name")}
        data.update(title="Эврика-2028", edition="XIII", year="2028", submissions_deadline="2028-03-01T18:00")
        data.pop("submissions_open", None)  # чекбокс снят — приём закрыт
        self.client.post(url, {**data, "application_pdf": SimpleUploadedFile("заявка.pdf", b"%PDF-1 old")})
        conf = ConferenceSettings.load()
        self.assertEqual((conf.title, conf.edition, conf.year), ("Эврика-2028", "XIII", 2028))
        self.assertFalse(conf.submissions_open)
        self.assertEqual(timezone.localtime(conf.submissions_deadline).hour, 18)
        old = Path(conf.application_pdf.path)
        self.assertTrue(old.exists())

        response = self.client.post(url, {**data, "application_pdf": SimpleUploadedFile("новая.pdf", b"%PDF-1 new")})
        self.assertRedirects(response, url)
        conf = ConferenceSettings.load()
        self.assertFalse(old.exists())  # старый бланк удалён с диска
        self.assertEqual(Path(conf.application_pdf.path).read_bytes(), b"%PDF-1 new")

        self.client.post(url, {**data, "application_pdf-clear": "on"})
        self.assertFalse(ConferenceSettings.load().application_pdf)

        response = self.client.post(url, {**data, "contact_email": "не почта"})
        self.assertContains(response, "Проверьте поля")

    def test_office_sections_crud(self):
        self.client.force_login(self.organizer)
        page = self.client.get(reverse("conference:office_site", args=["sections"]))
        self.assertContains(page, "Здоровое поколение")
        self.assertContains(page, "2/3 · 1/3")  # занято по секциям

        data = {"name": "Робототехника", "description": "", "groups": "2", "capacity": "10", "order": "5", "is_active": "on"}
        response = self.client.post(reverse("conference:office_site_add", args=["sections"]), data)
        self.assertRedirects(response, reverse("conference:office_site", args=["sections"]))
        new = Section.objects.get(name="Робототехника")
        self.assertEqual((new.groups, new.capacity), (2, 10))

        # Секций меньше, чем занято, — нельзя: во «Здоровом поколении 2» уже есть работа.
        edit_url = reverse("conference:office_site_edit", args=["sections", self.section.pk])
        response = self.client.post(edit_url, {"name": self.section.name, "groups": "1", "capacity": "3", "order": "0", "is_active": "on"})
        self.assertContains(response, "секций не может быть меньше 2")

        response = self.client.post(reverse("conference:office_site_delete", args=["sections", self.section.pk]))
        self.assertRedirects(response, edit_url)
        self.assertTrue(Section.objects.filter(pk=self.section.pk).exists())  # с работами не удаляется
        self.client.post(reverse("conference:office_site_delete", args=["sections", new.pk]))
        self.assertFalse(Section.objects.filter(pk=new.pk).exists())

    def test_office_news_and_documents(self):
        self.client.force_login(self.organizer)
        self.client.post(reverse("conference:office_site_add", args=["news"]), {
            "title": "Приём работ открыт", "text": "Ждём ваши работы.", "published_at": "2026-10-01T09:00", "is_published": "on",
        })
        self.assertContains(self.client.get(reverse("conference:home")), "Приём работ открыт")

        self.client.post(reverse("conference:office_site_add", args=["documents"]), {
            "title": "Положение", "description": "", "order": "1",
            "file": SimpleUploadedFile("положение.pdf", b"%PDF-1 old"),
        })
        document = Document.objects.get(title="Положение")
        self.assertFalse(document.is_published)  # галочка не стояла
        old = Path(document.file.path)
        # Скрытый документ скачивает оргкомитет, но не посетитель.
        response = self.client.get(reverse("conference:download_document", args=[document.pk]))
        self.assertEqual(response.status_code, 200)
        response.close()

        edit_url = reverse("conference:office_site_edit", args=["documents", document.pk])
        self.client.post(edit_url, {"title": "Положение", "description": "", "order": "1", "is_published": "on",
                                    "file": SimpleUploadedFile("новое.pdf", b"%PDF-1 new")})
        document.refresh_from_db()
        self.assertFalse(old.exists())
        self.assertTrue(document.is_published)
        new_path = Path(document.file.path)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("conference:office_site_delete", args=["documents", document.pk]))
        self.assertFalse(Document.objects.exists())
        self.assertFalse(new_path.exists())

    def test_news_cover_documents_and_page(self):
        from PIL import Image

        buffer = BytesIO()
        exif = Image.Exif()
        exif[0x8825] = {2: (55.0, 45.0, 0.0)}  # GPSInfo: место съёмки
        Image.new("RGB", (4000, 3000), (200, 30, 30)).save(buffer, "JPEG", exif=exif)

        self.client.force_login(self.organizer)
        data = {"title": "Конференция началась", "text": "Первый абзац.\n\nВторой абзац.",
                "published_at": "2026-10-01T09:00", "is_published": "on"}
        self.client.post(reverse("conference:office_site_add", args=["news"]), {
            **data,
            "cover": SimpleUploadedFile("обложка.jpg", buffer.getvalue()),
            "files": [SimpleUploadedFile("Программа.pdf", b"%PDF-1 program"), SimpleUploadedFile("Схема проезда.docx", b"docx")],
        })
        post = NewsPost.objects.get()
        with Image.open(post.cover.path) as image:
            self.assertEqual(image.size, (1600, 1200))
            self.assertFalse(image.getexif())  # без EXIF и GPS
        self.assertEqual([a.title for a in post.attachments.all()], ["Программа", "Схема проезда"])

        detail_url = reverse("conference:news_detail", args=[post.pk])
        home = self.client.get(reverse("conference:home"))
        self.assertContains(home, detail_url)
        self.assertContains(home, reverse("conference:news_cover", args=[post.pk]))
        self.assertNotContains(home, "Программа")  # документы — только на странице новости
        page = self.client.get(detail_url)
        self.assertContains(page, "Второй абзац.")
        program = post.attachments.get(title="Программа")
        self.assertContains(page, reverse("conference:news_file", args=[program.pk]))

        self.client.logout()
        response = self.client.get(reverse("conference:news_file", args=[program.pk]))
        from urllib.parse import quote

        self.assertIn(quote("Программа.pdf"), response["Content-Disposition"])
        response.close()
        response = self.client.get(reverse("conference:news_cover", args=[post.pk]))
        self.assertEqual(response["Content-Type"], "image/jpeg")
        response.close()  # иначе Windows не даст удалить открытый файл
        NewsPost.objects.update(is_published=False)
        for url in (detail_url, reverse("conference:news_cover", args=[post.pk]), reverse("conference:news_file", args=[program.pk])):
            self.assertEqual(self.client.get(url).status_code, 404)  # черновик скрыт

        self.client.force_login(self.organizer)
        self.assertContains(self.client.get(detail_url), "видит только оргкомитет")
        edit_url = reverse("conference:office_site_edit", args=["news", post.pk])
        response = self.client.post(edit_url, {**data, "cover": SimpleUploadedFile("текст.jpg", b"not an image")})
        self.assertContains(response, "не фото")
        response = self.client.post(edit_url, {**data, "files": [SimpleUploadedFile("вирус.exe", b"MZ")]})
        self.assertContains(response, "тип файла не подходит")

        cover, program_path = Path(post.cover.path), Path(program.file.path)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(edit_url, {**data, "cover-clear": "on", "remove_files": [program.pk]})
        post.refresh_from_db()
        self.assertFalse(post.cover)
        self.assertFalse(cover.exists())
        self.assertFalse(program_path.exists())
        other = Path(post.attachments.get().file.path)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("conference:office_site_delete", args=["news", post.pk]))
        self.assertFalse(other.exists())  # документы удаляются вместе с новостью

    def test_office_faq(self):
        self.assertEqual(FaqItem.objects.count(), 7)  # исправленные вопросы из миграции
        home = self.client.get(reverse("conference:home"))
        self.assertContains(home, "Можно ли исправить или удалить заявку?")
        self.assertContains(home, "СНИЛС соавтора")

        self.client.force_login(self.organizer)
        self.assertContains(self.client.get(reverse("conference:office_site", args=["faq"])), "Как понять, что работа дошла?")
        self.client.post(reverse("conference:office_site_add", args=["faq"]), {
            "question": "Где пройдёт конференция?", "answer": "В гимназии № 87.\n\nПодробности: https://example.ru", "order": "1",
            "is_published": "on",
        })
        item = FaqItem.objects.get(question="Где пройдёт конференция?")
        home = self.client.get(reverse("conference:home")).content.decode()
        self.assertLess(home.index("Где пройдёт конференция?"), home.index("Можно ли подать работу вдвоём?"))  # порядок 1 — первым
        self.assertIn('<a href="https://example.ru"', home)

        self.client.post(reverse("conference:office_site_edit", args=["faq", item.pk]),
                         {"question": item.question, "answer": item.answer, "order": "1"})  # снята галочка
        self.assertNotContains(self.client.get(reverse("conference:home")), "Где пройдёт конференция?")

    def test_office_search_by_snils_surname_and_number(self):
        self.client.force_login(self.organizer)
        target = Submission.objects.get(last_name="Фамилия1")
        digits = target.snils.replace("-", "").replace(" ", "")
        url = reverse("conference:office_search")
        self.assertContains(self.client.get(reverse("conference:office")), url)
        for query in (target.snils, digits, "фамилия1", str(target.pk)):
            response = self.client.get(url, {"q": query})
            self.assertEqual(list(response.context["results"]), [target], query)
        # СНИЛС соавтора находит его работу.
        Submission.objects.filter(pk=target.pk).update(coauthor_last_name="Соавтор", coauthor_first_name="Ира",
                                                        coauthor_grade=9, coauthor_snils="112-233-445 95")
        self.assertEqual(list(self.client.get(url, {"q": "11223344595"}).context["results"]), [target])
        self.assertContains(self.client.get(url, {"q": "Нетакого"}), "Ничего не найдено")

    def test_office_move_and_edit_submission(self):
        self.client.force_login(self.organizer)
        submission = Submission.objects.get(last_name="Фамилия0")  # «Здоровое поколение 1», capacity 3
        url = reverse("conference:office_submission", args=[submission.pk])
        page = self.client.get(url)
        self.assertContains(page, "Здоровое поколение 2 — занято 1 из 3")
        self.assertContains(page, "Краеведение — занято 0 из 12")

        response = self.client.post(url, {"action": "move", "slot": f"{self.other.pk}:1"})
        self.assertRedirects(response, url)
        submission.refresh_from_db()
        self.assertEqual((submission.section, submission.section_number), (self.other, 1))

        # Секция заполнена: перенести нельзя, место не создаётся.
        self.other.capacity = 1
        self.other.save()
        Submission.objects.filter(last_name="Фамилия2").update(section=self.other, section_number=1)
        mover = Submission.objects.get(last_name="Фамилия1")
        response = self.client.post(reverse("conference:office_submission", args=[mover.pk]),
                                    {"action": "move", "slot": f"{self.other.pk}:1"}, follow=True)
        self.assertContains(response, "нет свободных мест")
        mover.refresh_from_db()
        self.assertNotEqual(mover.section, self.other)

        data = {name: value for name, value in page.context["form"].initial.items()
                if value is not None and name not in ("file", "application_file", "consent_file")}
        data.update(last_name="Иванова", snils="112-233-445 95", school="Лицей № 1")
        response = self.client.post(url, {**data, "file": SimpleUploadedFile("новая.docx", b"new")})
        self.assertRedirects(response, url)
        submission.refresh_from_db()
        self.assertEqual((submission.last_name, submission.school, submission.original_filename), ("Иванова", "Лицей № 1", "новая.docx"))

        response = self.client.post(url, {**data, "snils": "112-233-445 96"})
        self.assertContains(response, "опечатка")  # та же проверка СНИЛС, что в анкете

    def test_office_submission_edit_requires_change_permission(self):
        viewer = User.objects.create_user("viewer", password="Pass-1234!", is_staff=True)
        viewer.user_permissions.add(Permission.objects.get(codename="view_submission"))
        self.client.force_login(viewer)
        submission = Submission.objects.first()
        self.assertEqual(self.client.get(reverse("conference:office_submission", args=[submission.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse("conference:office_search"), {"q": submission.last_name}).status_code, 200)

    def test_office_site_pages_require_permissions(self):
        viewer = User.objects.create_user("viewer", password="Pass-1234!", is_staff=True)
        viewer.user_permissions.add(Permission.objects.get(codename="view_submission"))
        self.client.force_login(viewer)
        for kind in ("sections", "news", "faq", "documents"):
            self.assertEqual(self.client.get(reverse("conference:office_site", args=[kind])).status_code, 403)
        self.assertEqual(self.client.get(reverse("conference:office_site", args=["users"])).status_code, 404)

    def test_settings_page_requires_permission(self):
        viewer = User.objects.create_user("viewer", password="Pass-1234!", is_staff=True)
        viewer.user_permissions.add(Permission.objects.get(codename="view_submission"))
        self.client.force_login(viewer)
        self.assertNotContains(self.client.get(reverse("conference:office")), reverse("conference:office_settings"))
        self.assertEqual(self.client.get(reverse("conference:office_settings")).status_code, 403)

    def test_admission_requires_change_permission(self):
        viewer = User.objects.create_user("viewer", password="Pass-1234!", is_staff=True)
        viewer.user_permissions.add(Permission.objects.get(codename="view_submission"))
        self.client.force_login(viewer)
        conf = ConferenceSettings.load()
        conf.submissions_open = False
        conf.save()
        submission = Submission.objects.get(last_name="Фамилия0")
        response = self.client.post(reverse("conference:office_admission", args=[submission.pk]), {"admission": "admitted"})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.post(reverse("conference:office_intake"), {"open": "1"}).status_code, 403)

    def test_section_page_and_exports(self):
        from docx import Document
        from openpyxl import load_workbook

        self.client.force_login(self.organizer)
        response = self.client.get(reverse("conference:office_section", args=[self.section.pk, 1]))
        self.assertContains(response, "Фамилия0 Имя")
        self.assertNotContains(response, "Фамилия2 Имя")  # она во второй секции

        docx = self.client.get(reverse("conference:office_section_word", args=[self.section.pk, 1]))
        text = "\n".join(cell.text for t in Document(BytesIO(docx.getvalue())).tables for row in t.rows for cell in row.cells)
        self.assertIn("Фамилия1 Имя", text)
        self.assertEqual(self.client.get(reverse("conference:office_word")).status_code, 200)

        wb = load_workbook(BytesIO(b"".join(self.client.get(reverse("conference:office_participants")).streaming_content)))
        names = [row[3] for row in wb.active.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(names), ["Фамилия0 Имя", "Фамилия1 Имя", "Фамилия2 Имя"])

    def test_office_dashboard_stage_one_exports(self):
        from docx import Document
        from openpyxl import load_workbook

        self.client.force_login(self.organizer)
        dashboard = self.client.get(reverse("conference:office"))
        self.assertContains(dashboard, reverse("conference:office_students_excel"))
        self.assertContains(dashboard, reverse("conference:office_students_word"))
        response = self.client.get(reverse("conference:office_students_excel"))
        rows = list(load_workbook(BytesIO(b"".join(response.streaming_content))).active.values)
        self.assertEqual(len(rows), 2 + 1 + 3)  # заголовок, пустая строка, шапка, ученики
        response = self.client.get(reverse("conference:office_students_word"))
        self.assertEqual(len(Document(BytesIO(response.getvalue())).tables[0].rows), 1 + 3)

    def test_admitted_exports_only_admitted_sorted_by_section(self):
        from docx import Document
        from openpyxl import load_workbook

        other = Section.objects.create(name="Абвгд", order=99)
        Submission.objects.filter(last_name="Фамилия0").update(admission="admitted")
        Submission.objects.filter(last_name="Фамилия1").update(admission="not_admitted")
        Submission.objects.create(
            section=other, admission="admitted", snils=make_snils(400000001), last_name="Яблоков", first_name="Пётр",
            grade=10, email="y@example.ru", school="Школа", district="Р", title="Работа Я",
            supervisor_last_name="Иванова", supervisor_first_name="Мария", supervisor_position="учитель",
            supervisor_phone="+79001234567", file="x.docx",
        )
        self.client.force_login(self.organizer)
        response = self.client.get(reverse("conference:office_admitted_excel"))
        rows = list(load_workbook(BytesIO(b"".join(response.streaming_content))).active.values)
        self.assertIn("ко 2 этапу XII межрегиональной", rows[0][0])
        self.assertIn("в 2027 году", rows[0][0])
        conf = ConferenceSettings.load()
        conf.edition, conf.year = "XIII", 2028
        conf.save()
        response = self.client.get(reverse("conference:office_admitted_excel"))
        title = list(load_workbook(BytesIO(b"".join(response.streaming_content))).active.values)[0][0]
        self.assertIn("ко 2 этапу XIII межрегиональной", title)
        self.assertIn("в 2028 году", title)
        data = rows[3:]
        self.assertEqual([r[2] for r in data], ["Яблоков Пётр", "Фамилия0 Имя"])  # «Абвгд» раньше по алфавиту
        self.assertEqual(len(data[0]), 7)  # без даты и почты

        word = self.client.get(reverse("conference:office_admitted_word"))
        self.assertEqual(len(Document(BytesIO(word.getvalue())).tables[0].rows), 1 + 2)
        pdf = self.client.get(reverse("conference:office_admitted_pdf"))
        self.assertTrue(pdf.getvalue().startswith(b"%PDF"))

    def test_custom_export_formats_fields_and_filters(self):
        from docx import Document
        from openpyxl import load_workbook

        self.client.force_login(self.organizer)
        base = {
            "field": ["number", "name", "title"],
            "group": [f"{self.section.pk}:1"],
            "admission": ["pending", "admitted", "not_admitted"],
        }
        url = reverse("conference:office_custom")

        response = self.client.get(url, {**base, "format": "excel", "title": "Мой список", "rows": "student", "sort": "alpha"})
        rows = list(load_workbook(BytesIO(b"".join(response.streaming_content))).active.values)
        self.assertEqual(rows[0][0], "Мой список")
        self.assertEqual(rows[2], ("№ п/п", "ФИО", "Название работы"))
        self.assertTrue(all(len(r) == 3 for r in rows[3:]))
        self.assertNotIn("Фамилия2 Имя", [r[1] for r in rows[3:]])  # вторая секция не отмечена

        word = self.client.get(url, {**base, "format": "word"})
        self.assertGreater(len(Document(BytesIO(word.getvalue())).tables[0].rows), 1)
        pdf = self.client.get(url, {**base, "format": "pdf"})
        self.assertTrue(pdf.getvalue().startswith(b"%PDF"))

        for broken in ({**base, "field": []}, {**base, "group": []}, {**base, "admission": []}, {**base, "format": "exe"}):
            self.assertEqual(self.client.get(url, broken).status_code, 302)

    def test_custom_export_requires_office_login(self):
        response = self.client.get(reverse("conference:office_custom"), {"field": "name"})
        self.assertEqual(response.status_code, 302)
        self.assertIn("login", response["Location"])

    def test_action_log_records_decisions_and_moves(self):
        from conference.models import ActionLog

        self.client.force_login(self.organizer)
        Submission.objects.update(admission="")
        conf = ConferenceSettings.load()
        conf.submissions_open = False
        conf.save()
        submission = Submission.objects.get(last_name="Фамилия0")

        self.client.post(reverse("conference:office_admission", args=[submission.pk]), {"admission": "admitted"})
        self.client.post(reverse("conference:office_award", args=[submission.pk]), {"award": "place1"})
        self.client.post(reverse("conference:office_admission", args=[submission.pk]), {"admission": "not_admitted"})
        self.client.post(reverse("conference:office_submission", args=[submission.pk]), {"action": "move", "slot": f"{self.other.pk}:1"})

        entries = list(ActionLog.objects.order_by("pk"))
        self.assertEqual([e.action for e in entries], ["admission", "award", "admission", "move"])
        self.assertIn("не решено → допущен", entries[0].details)
        self.assertIn("нет → 1 место", entries[1].details)
        self.assertIn("Итог «1 место» снят", entries[2].details)
        self.assertEqual(entries[0].user_name, self.organizer.username)

        page = self.client.get(reverse("conference:office_log"), {"s": submission.pk})
        self.assertContains(page, "→")
        self.assertEqual(len(page.context["entries"]), 4)

    def test_model_blocks_full_section_and_lower_capacity(self):
        from django.core.exceptions import ValidationError

        full = Section.objects.create(name="Полная", capacity=1)
        first = Submission.objects.get(last_name="Фамилия0")
        Submission.objects.filter(pk=first.pk).update(section=full, section_number=1)
        second = Submission.objects.get(last_name="Фамилия1")
        second.section, second.section_number = full, 1
        with self.assertRaises(ValidationError):
            second.clean()
        first.refresh_from_db()
        first.clean()  # уже занятое место отнять нельзя
        full.capacity = 5
        full.save()
        full.capacity = 0
        with self.assertRaises(ValidationError):
            full.clean()

    def test_archive_contains_works_and_list(self):
        import zipfile

        self.client.force_login(self.organizer)
        response = self.client.get(reverse("conference:office_archive"))
        self.assertIn("attachment", response["Content-Disposition"])
        archive = zipfile.ZipFile(BytesIO(b"".join(response.streaming_content)))
        names = archive.namelist()
        self.assertIn("Список работ.xlsx", names)
        works = [n for n in names if n.endswith(".docx")]
        self.assertEqual(len(works), 3)
        self.assertTrue(any(n.startswith("Здоровое поколение 2/") for n in works))
        first = next(n for n in works if "Фамилия0" in n)
        self.assertEqual(archive.read(first).decode(), "содержимое 0")

    def test_review_link_flow(self):
        import zipfile

        self.client.force_login(self.organizer)
        response = self.client.post(reverse("conference:office_link_create"), {"section": self.section.pk, "number": 1})
        link = ReviewLink.objects.get()
        self.assertRedirects(response, reverse("conference:office_link", args=[link.pk]))
        # повторное нажатие не создаёт вторую ссылку
        self.client.post(reverse("conference:office_link_create"), {"section": self.section.pk, "number": 1})
        self.assertEqual(ReviewLink.objects.count(), 1)

        self.client.logout()
        page = self.client.get(reverse("conference:review", args=[link.token]))
        self.assertContains(page, "2 работы")
        archive = zipfile.ZipFile(BytesIO(b"".join(self.client.get(reverse("conference:review_download", args=[link.token])).streaming_content)))
        self.assertEqual(len([n for n in archive.namelist() if n.endswith(".docx")]), 2)
        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(archive.read("Список работ.xlsx"))).active
        self.assertNotIn("p0@example.ru", str(list(sheet.values)))  # без контактов

        link.is_active = False
        link.save()
        self.assertEqual(self.client.get(reverse("conference:review", args=[link.token])).status_code, 410)
        self.assertEqual(self.client.get(reverse("conference:review_download", args=[link.token])).status_code, 404)

    def test_expired_link_does_not_work(self):
        link = ReviewLink.objects.create(expires_at=timezone.now() - timedelta(minutes=1))
        self.assertEqual(self.client.get(reverse("conference:review_download", args=[link.token])).status_code, 404)
        self.assertEqual(ReviewLink.objects.create().expires_at.date(), (timezone.now() + timedelta(days=30)).date())


class OfficePasswordTests(TestCase):
    def setUp(self):
        self.organizer = User.objects.create_user("olga", password="Старый-пароль-1", is_staff=True, first_name="Ольга")
        self.organizer.groups.add(Group.objects.get(name=ORGANIZERS_GROUP))
        self.admin = User.objects.create_superuser("adminK", password="Админ-пароль-1")

    def test_staff_changes_own_password_and_stays_logged_in(self):
        self.client.force_login(self.organizer)
        response = self.client.post(reverse("conference:office_password"), {
            "old_password": "Старый-пароль-1", "new_password1": "сова-трамвай-лимон-7", "new_password2": "сова-трамвай-лимон-7",
        })
        self.assertRedirects(response, reverse("conference:office"))
        self.organizer.refresh_from_db()
        self.assertTrue(self.organizer.check_password("сова-трамвай-лимон-7"))
        self.assertEqual(self.client.get(reverse("conference:office")).status_code, 200)

    def test_wrong_old_password_and_weak_password_rejected(self):
        self.client.force_login(self.organizer)
        response = self.client.post(reverse("conference:office_password"), {
            "old_password": "не тот", "new_password1": "123456", "new_password2": "123456",
        })
        self.assertEqual(response.status_code, 200)
        self.organizer.refresh_from_db()
        self.assertTrue(self.organizer.check_password("Старый-пароль-1"))

    def test_superuser_sets_password_for_staff(self):
        self.client.force_login(self.admin)
        self.assertContains(self.client.get(reverse("conference:office_staff")), "olga")
        response = self.client.post(reverse("conference:office_staff_password", args=[self.organizer.pk]), {
            "new_password1": "сова-трамвай-лимон-7", "new_password2": "сова-трамвай-лимон-7",
        })
        self.assertRedirects(response, reverse("conference:office_staff"))
        self.organizer.refresh_from_db()
        self.assertTrue(self.organizer.check_password("сова-трамвай-лимон-7"))

    def test_organizer_cannot_manage_staff(self):
        self.client.force_login(self.organizer)
        self.assertEqual(self.client.get(reverse("conference:office_staff")).status_code, 403)
        url = reverse("conference:office_staff_password", args=[self.admin.pk])
        self.assertEqual(self.client.post(url, {"new_password1": "x-y-z-12345", "new_password2": "x-y-z-12345"}).status_code, 403)


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class TestSubmissionsCommandTests(TestCase):
    def test_create_and_delete_keep_real_submissions(self):
        from django.core.management import call_command

        section = Section.objects.create(name="Здоровое поколение", groups=2)
        real = Submission.objects.create(
            section=section, snils=make_snils(500000001), last_name="Настоящая", first_name="Анна", grade=9,
            email="real@example.ru", school="Школа", district="Район", title="Т",
            supervisor_last_name="П", supervisor_first_name="А", supervisor_position="учитель",
            supervisor_phone="+79001234567", file="x.docx",
        )
        call_command("test_submissions", stdout=StringIO())
        tests = Submission.objects.exclude(pk=real.pk)
        self.assertTrue(tests.exists())
        self.assertTrue(all(s.email.endswith("@eureka-test.example") for s in tests))
        self.assertLessEqual(section.occupied()[1], section.capacity)
        call_command("test_submissions", "--delete", stdout=StringIO())
        self.assertEqual(list(Submission.objects.all()), [real])
