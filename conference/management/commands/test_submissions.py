"""Тестовые заявки для обучения оргкомитета и проверки выгрузок.

    manage.py test_submissions            создать (если тестовых ещё нет)
    manage.py test_submissions --delete   удалить все тестовые заявки вместе с файлами

Тестовые заявки помечены почтой на домене @eureka-test.example и СНИЛС из служебного
диапазона (номера до 001-001-998 не выдаются людям), поэтому не пересекаются с настоящими.
"""

import random
from datetime import timedelta
from io import BytesIO

from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from docx import Document

from conference.models import Section, Submission

TEST_DOMAIN = "eureka-test.example"

MALE = ["Александр", "Дмитрий", "Максим", "Иван", "Артём", "Никита", "Михаил", "Егор", "Кирилл", "Андрей",
        "Илья", "Матвей", "Тимофей", "Роман", "Даниил", "Арсений", "Глеб", "Степан"]
FEMALE = ["Анна", "Мария", "Дарья", "Полина", "Елизавета", "Виктория", "Анастасия", "Софья", "Алиса",
          "Варвара", "Ксения", "Вероника", "Екатерина", "Ульяна", "Милана", "Арина", "Ева", "Таисия"]
SURNAMES = ["Иванов", "Смирнов", "Кузнецов", "Попов", "Васильев", "Петров", "Соколов", "Михайлов", "Новиков",
            "Фёдоров", "Морозов", "Волков", "Алексеев", "Лебедев", "Семёнов", "Егоров", "Павлов", "Козлов",
            "Степанов", "Николаев", "Орлов", "Андреев", "Макаров", "Никитин", "Захаров", "Зайцев", "Соловьёв",
            "Борисов", "Яковлев", "Григорьев", "Романов", "Воробьёв", "Сергеев", "Фролов", "Белов", "Тарасов"]
MIDDLE_F = ["Александровна", "Сергеевна", "Викторовна", "Николаевна", "Владимировна", "Андреевна", "Игоревна",
            "Петровна", "Михайловна", "Олеговна"]
SCHOOLS = [
    ("Муниципальное автономное общеобразовательное учреждение «Гимназия № 87» Ленинского района г. Саратова",
     "Ленинский район г. Саратова"),
    ("Муниципальное общеобразовательное учреждение «Средняя общеобразовательная школа № 12» Кировского района г. Саратова",
     "Кировский район г. Саратова"),
    ("Муниципальное автономное общеобразовательное учреждение «Лицей № 2» Фрунзенского района г. Саратова",
     "Фрунзенский район г. Саратова"),
    ("Муниципальное общеобразовательное учреждение «Средняя общеобразовательная школа № 45» Волжского района г. Саратова",
     "Волжский район г. Саратова"),
    ("Муниципальное автономное общеобразовательное учреждение «Лицей № 36» Октябрьского района г. Саратова",
     "Октябрьский район г. Саратова"),
    ("Муниципальное общеобразовательное учреждение «Средняя общеобразовательная школа № 9» Заводского района г. Саратова",
     "Заводской район г. Саратова"),
    ("Муниципальное общеобразовательное учреждение «Средняя общеобразовательная школа № 32» г. Энгельса",
     "Энгельсский муниципальный район"),
    ("Муниципальное общеобразовательное учреждение «Средняя общеобразовательная школа № 3» г. Балаково",
     "Балаковский муниципальный район"),
]
POSITIONS = ["учитель биологии", "учитель истории", "учитель русского языка и литературы", "учитель математики",
             "учитель информатики", "учитель географии", "учитель английского языка", "учитель химии",
             "педагог дополнительного образования", "учитель физики"]

# Темы работ по ключевому слову в названии направления.
TOPICS = {
    "Навыки": ["Тайм-менеджмент школьника: как успеть всё", "Критическое мышление и фейковые новости",
               "Финансовая грамотность подростка", "Командная работа в школьных проектах"],
    "Наследие": ["Судьба моего прадеда в годы войны", "Забытые ремёсла Саратовского края",
                 "Памятники нашего района: история и современность", "Семейный архив как исторический источник"],
    "Семья": ["Семейные традиции в современном мире", "Роль бабушек и дедушек в воспитании",
              "Как цифровые устройства меняют семейное общение", "Семейный бюджет глазами подростка"],
    "Здоровое": ["Влияние режима сна на успеваемость", "Энергетические напитки: польза или вред",
                 "Осанка и школьный рюкзак", "Сколько сахара в школьном обеде", "Экранное время и зрение"],
    "Искусство слова": ["Неологизмы в речи подростков", "Образ города в поэзии XX века",
                        "Язык мемов как явление современной речи", "Анализ рекламных слоганов"],
    "краеведение": ["История улицы, на которой я живу", "Купеческий Саратов: особняки и их владельцы",
                    "Родники Саратовской области", "Музей моей школы", "Волга в жизни саратовцев"],
    "Экология": ["Раздельный сбор мусора в нашей школе", "Качество воды в реке Волге",
                 "Лишайники как индикатор чистоты воздуха", "Пластик в нашей жизни", "Птицы городского парка"],
    "Языкознание": ["Англицизмы в молодёжном сленге", "Британские и американские традиции Рождества",
                    "Немецкие заимствования в русском языке", "Фразеологизмы о еде в разных языках"],
    "Азбучные": ["Чат-бот для расписания уроков", "Алгоритмы сортировки: что быстрее",
                 "Безопасность паролей школьников", "Как устроен поиск в интернете"],
    "Ракурс": ["Золотое сечение в архитектуре Саратова", "Математика в шахматах",
               "Теория вероятностей в настольных играх", "Симметрия в природе"],
    "географии": ["Топонимика Саратовской области", "Климат Саратова за 50 лет",
                  "Туристический маршрут по родному краю", "Почему Волга меняет русло"],
    "Трудные": ["Саратов в годы Великой Отечественной войны", "Реформы Петра I: взгляд современника",
                "Эвакуация предприятий в Саратов", "Смутное время в учебниках разных лет"],
    "природы": ["Выращивание кристаллов в домашних условиях", "Как растения реагируют на музыку",
                "Микроорганизмы на школьных поверхностях", "Маятник Фуко в домашних условиях"],
}

# Сколько мест занять в секциях направления (по порядку): разные случаи для проверки выгрузок.
FILL_PATTERNS = [
    [12, 10, 3],   # первая заполнена, вторая почти, третья только начата
    [12, 6],
    [9],
    [11],
    [12, 12, 5],   # две полные секции
    [7, 0],
    [4],
    [12],          # заполнена целиком
    [0],           # пустая
    [8, 0, 0],
]

# Маленький корректный PDF для «скана» заявки.
PDF = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
       b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 100]>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF\n")
# PNG 1×1 для «фото» согласия.
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360f8cf00000301010018dd8db00000000049454e44ae426082"
)


def test_snils(n):
    """СНИЛС из служебного диапазона (до 001-001-998): контрольное число не проверяется, людям не выдаётся."""
    digits = f"{n:09d}"
    return f"{digits[:3]}-{digits[3:6]}-{digits[6:]} 00"


def topics_for(section):
    for key, items in TOPICS.items():
        if key.lower() in section.name.lower():
            return items
    return ["Исследовательская работа по направлению «%s»" % section.name]


def work_docx(submission):
    doc = Document()
    doc.add_paragraph("ТЕСТОВАЯ РАБОТА — создана для обучения оргкомитета").runs[0].bold = True
    doc.add_heading(submission.title, level=1)
    doc.add_paragraph(f"Автор: {submission.full_name}, {submission.grade} класс")
    doc.add_paragraph(f"Образовательная организация: {submission.school}")
    doc.add_paragraph(f"Руководитель: {submission.supervisor}, {submission.supervisor_position}")
    doc.add_heading("Введение", level=2)
    doc.add_paragraph("Актуальность исследования. Цель и задачи работы. " * 6)
    doc.add_heading("Заключение", level=2)
    doc.add_paragraph("Выводы по результатам исследования. " * 5)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class Command(BaseCommand):
    help = "Создаёт тестовые заявки для обучения оргкомитета (или удаляет их с --delete)."

    def add_arguments(self, parser):
        parser.add_argument("--delete", action="store_true", help="удалить все тестовые заявки и их файлы")

    def handle(self, *args, delete=False, **options):
        tests = Submission.objects.filter(email__endswith="@" + TEST_DOMAIN)
        if delete:
            count = 0
            for s in tests:
                for f in (s.file, s.application_file, s.consent_file):
                    if f:
                        f.delete(save=False)
                s.delete()
                count += 1
            self.stdout.write(self.style.SUCCESS(f"Удалено тестовых заявок: {count}."))
            return

        if tests.exists():
            raise CommandError(
                f"Тестовые заявки уже есть ({tests.count()}). Сначала удалите их: manage.py test_submissions --delete"
            )
        sections = list(Section.objects.filter(is_active=True))
        if not sections:
            raise CommandError("Нет направлений. Сначала: manage.py load_sections")

        rng = random.Random(2027)
        created = 0
        serial = 1000
        now = timezone.now()
        with transaction.atomic():
            for index, section in enumerate(sections):
                pattern = FILL_PATTERNS[index % len(FILL_PATTERNS)]
                topics = topics_for(section)
                for number in range(1, section.groups + 1):
                    target = min(pattern[number - 1] if number - 1 < len(pattern) else 0, section.capacity)
                    free = section.free_places().get(number, 0)
                    target = min(target, free)
                    people = 0
                    while people < target:
                        with_coauthor = target - people >= 2 and rng.random() < 0.18
                        serial += 1
                        s = self._make(rng, section, number, topics, serial, with_coauthor)
                        s.save()
                        # Даты подачи — за последние три недели, чтобы список выглядел живым.
                        created_at = now - timedelta(days=rng.randint(0, 20), minutes=rng.randint(0, 1400))
                        Submission.objects.filter(pk=s.pk).update(created_at=created_at)
                        if s.status != Submission.Status.REJECTED:  # отклонённые места не занимают
                            people += s.people
                        created += 1
        self.stdout.write(self.style.SUCCESS(f"Создано тестовых заявок: {created}."))
        self.stdout.write("Удалить их: manage.py test_submissions --delete")

    def _person(self, rng):
        female = rng.random() < 0.55
        first = rng.choice(FEMALE if female else MALE)
        last = rng.choice(SURNAMES)
        if female:
            last += "а"
        return last, first

    def _make(self, rng, section, number, topics, serial, with_coauthor):
        last, first = self._person(rng)
        school, district = rng.choice(SCHOOLS)
        sup_last = rng.choice(SURNAMES) + "а"
        title = rng.choice(topics)
        status = rng.choices(["new", "revision", "rejected", "accepted"], weights=[80, 8, 4, 8])[0]
        s = Submission(
            section=section,
            section_number=number,
            status=status,
            last_name=last,
            first_name=first,
            grade=rng.randint(8, 11),
            snils=test_snils(serial),
            email=f"pupil{serial}@{TEST_DOMAIN}",
            school=school,
            district=district,
            title=title,
            supervisor_last_name=sup_last,
            supervisor_first_name=rng.choice(FEMALE),
            supervisor_middle_name=rng.choice(MIDDLE_F),
            supervisor_position=rng.choice(POSITIONS),
            supervisor_phone=f"+7 (9{rng.randint(10, 99)}) {rng.randint(100, 999)}-{rng.randint(10, 99)}-{rng.randint(10, 99)}",
        )
        if status == "revision":
            s.participant_message = "Добавьте список литературы и оформите титульный лист по требованиям."
        if with_coauthor:
            c_last, c_first = self._person(rng)
            s.coauthor_last_name, s.coauthor_first_name = c_last, c_first
            s.coauthor_grade = s.grade
            s.coauthor_snils = test_snils(serial + 500000)
        if rng.random() < 0.12:
            s.supervisor2_last_name = rng.choice(SURNAMES) + "а"
            s.supervisor2_first_name = rng.choice(FEMALE)
            s.supervisor2_middle_name = rng.choice(MIDDLE_F)
            s.supervisor2_position = rng.choice(POSITIONS)
            s.supervisor2_phone = f"+7 (9{rng.randint(10, 99)}) {rng.randint(100, 999)}-00-00"
        s.file.save(f"{title}.docx", ContentFile(work_docx(s)), save=False)
        s.original_filename = f"{title}.docx"
        s.application_file.save("заявка.pdf", ContentFile(PDF), save=False)
        s.application_original_filename = "заявка.pdf"
        s.consent_file.save("согласие.png", ContentFile(PNG), save=False)
        s.consent_original_filename = "согласие.png"
        return s
